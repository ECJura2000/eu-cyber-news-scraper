from __future__ import annotations

import json
import math
import os
import platform
import subprocess
import tempfile
from datetime import datetime, timedelta
from hashlib import sha256
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import default_sources_path
from .coverage import load_coverage
from .models import Article, Source, SourceStatus
from .periods import PeriodSelection
from .translation import TranslationReport

ARTICLE_HEADERS = (
    "編號",
    "國家／區域",
    "發布機關",
    "機關屬性",
    "規範權威層級",
    "內容性質",
    "發布日期",
    "原始日期文字",
    "來源時區",
    "當地發布日",
    "日期精度",
    "日期來源",
    "日期信心",
    "日期衝突",
    "原文語言",
    "新聞標題",
    "繁體中文標題",
    "摘要",
    "觀測主題",
    "命中關鍵字",
    "關聯分數",
    "可信度",
    "官方原文",
    "其他官方連結",
    "抓取方式",
    "發現來源",
)

FORMULA_PREFIXES = ("=", "+", "-", "@")
HIGH_CONFIDENCE_FILL = PatternFill("solid", fgColor="FFC000")
MEDIUM_CONFIDENCE_FILL = PatternFill("solid", fgColor="FFF2CC")
LOW_CONFIDENCE_FILL = PatternFill("solid", fgColor="FFFCE6")
EXCEL_SUMMARY_LIMIT = 500


def safe_excel_text(value: str) -> str:
    """Prevent externally supplied text from being interpreted as a formula."""
    if value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def export_workbook(
    articles: list[Article],
    statuses: list[SourceStatus],
    sources: list[Source],
    output_path: str | Path,
) -> Path:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "全部命中新聞"
    _write_articles(ws_all, articles, highlight_matches=True)

    ws_law = wb.create_sheet("CRA_CSA_NIS2_CER")
    legal_topics = {
        "產品資安、漏洞揭露義務、SBOM",
        "資安產品認證",
        "NIS2、關鍵基礎設施保護",
        "關鍵基礎設施的實體與數位融合保護",
    }
    _write_articles(ws_law, [item for item in articles if legal_topics.intersection(item.matched_topics)])

    ws_authority = wb.create_sheet("官方規範與執法")
    _write_articles(
        ws_authority,
        [item for item in articles if item.authority_level in {"立法機關", "主管／監理機關", "法定機構／行政法人"}],
    )

    ws_research = wb.create_sheet("研究智庫與公私協力")
    _write_articles(
        ws_research,
        [item for item in articles if item.authority_level in {"公私協力／產學平台", "政策智庫", "研究／支援機構"}],
    )

    ws_status = wb.create_sheet("來源健康狀態")
    _write_statuses(ws_status, statuses)

    ws_sources = wb.create_sheet("官方來源清單")
    _write_sources(ws_sources, sources)

    ws_coverage = wb.create_sheet("議題主管機關覆蓋")
    _write_coverage(ws_coverage, sources)

    temporary = path.with_name(f"{path.stem}.tmp{path.suffix}")
    try:
        wb.save(temporary)
        _verify_workbook(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)
    return path


def write_jsonl(articles: list[Article], output_path: str | Path, *, run_id: str = "") -> Path:
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = []
    for article in articles:
        lines.append(
            json.dumps(
                {
                    "schema_version": 5,
                    "run_id": run_id,
                    "source_id": article.source_id,
                    "country": article.country,
                    "source_name": article.source_name,
                    "institution_type": article.institution_type,
                    "language": article.language,
                    "published_at": article.published_at.isoformat() if article.published_at else None,
                    "published_at_raw": article.published_at_raw,
                    "published_date_local": article.published_date_local,
                    "published_timezone": article.published_timezone,
                    "date_precision": article.date_precision,
                    "date_source": article.date_source,
                    "date_confidence": article.date_confidence,
                    "date_conflict": article.date_conflict,
                    "title": article.title,
                    "title_zh_tw": article.title_zh_tw,
                    "summary": article.summary,
                    "url": article.url,
                    "alternate_urls": article.alternate_urls,
                    "matched_topics": article.matched_topics,
                    "matched_keywords": article.matched_keywords,
                    "relevance_score": article.relevance_score,
                    "confidence_level": article.confidence_level,
                    "fetched_via": article.fetched_via,
                    "discovered_by": article.discovered_by,
                    "authority_level": article.authority_level,
                    "content_kind": article.content_kind,
                },
                ensure_ascii=False,
            )
        )
    _atomic_write_text(path, "\n".join(lines) + ("\n" if lines else ""))
    return path


def _write_articles(ws: Any, articles: list[Article], *, highlight_matches: bool = False) -> None:
    ws.append(ARTICLE_HEADERS)
    for index, article in enumerate(articles, 1):
        ws.append(
            (
                index,
                article.country,
                safe_excel_text(article.source_name),
                safe_excel_text(article.institution_type),
                article.authority_level,
                article.content_kind,
                article.published_at.date().isoformat() if article.published_at else "未辨識",
                safe_excel_text(article.published_at_raw),
                article.published_timezone,
                article.published_date_local,
                article.date_precision,
                article.date_source,
                article.date_confidence,
                "是" if article.date_conflict else "否",
                article.language,
                safe_excel_text(article.title),
                safe_excel_text(article.title_zh_tw),
                safe_excel_text(_excel_summary(article.summary)),
                "；".join(article.matched_topics),
                "；".join(article.matched_keywords),
                article.relevance_score,
                article.confidence_level,
                article.url,
                "\n".join(article.alternate_urls),
                article.fetched_via,
                "；".join(article.discovered_by or [article.source_id]),
            )
        )
        ws.cell(ws.max_row, 23).hyperlink = article.url
        ws.cell(ws.max_row, 23).style = "Hyperlink"
        if highlight_matches and article.matched_keywords:
            fill = {
                "高": HIGH_CONFIDENCE_FILL,
                "中": MEDIUM_CONFIDENCE_FILL,
                "低": LOW_CONFIDENCE_FILL,
            }.get(article.confidence_level, LOW_CONFIDENCE_FILL)
            for cell in ws[ws.max_row]:
                cell.fill = fill
    _style_table(
        ws,
        widths=(8, 12, 28, 18, 18, 18, 18, 30, 20, 18, 14, 20, 14, 12, 12, 56, 56, 72, 36, 44, 12, 12, 64, 64, 28, 32),
    )


def _write_statuses(ws: Any, statuses: list[SourceStatus]) -> None:
    ws.append(("來源代碼", "國家", "機關", "必要來源", "抓取狀態", "解析狀態", "新鮮度狀態", "內容狀態", "健康狀態", "抓取方式", "抓取頁數", "HTTP 請求", "下載位元組", "重試次數", "HTTP 狀態碼", "預算耗盡", "原始筆數", "期間內筆數", "有日期筆數", "無效未來日期", "逾時次數", "無日期比例", "唯一標題比例", "歷史中位數", "命中筆數", "秒數", "最新日期", "新鮮度落後天數", "健康警示", "抓取警示", "內容結果", "錯誤"))
    for status in statuses:
        ws.append(
            (
                status.source_id,
                status.country,
                safe_excel_text(status.source_name),
                "是" if status.critical else "否",
                status.fetch_status,
                status.parse_status,
                status.freshness_status,
                status.content_status,
                status.health_status,
                status.fetched_via,
                status.pages_fetched,
                status.request_count,
                status.bytes_downloaded,
                status.retry_count,
                "、".join(status.http_statuses),
                "是" if status.budget_exhausted else "否",
                status.raw_count,
                status.in_range_count,
                status.dated_count,
                status.invalid_date_count,
                status.timeout_count,
                status.undated_ratio,
                status.unique_title_ratio,
                status.historical_median_count,
                status.relevant_count,
                status.duration_seconds,
                status.newest_published_at,
                status.freshness_lag_days,
                " ".join(status.health_alerts),
                safe_excel_text(status.warning),
                safe_excel_text(status.content_warning),
                safe_excel_text(status.error),
            )
        )
    _style_table(ws, widths=(22, 10, 30, 12, 14, 14, 14, 14, 14, 20, 12, 12, 16, 12, 24, 12, 12, 12, 12, 16, 12, 14, 14, 14, 12, 10, 26, 16, 52, 52, 52, 72))


def _write_sources(ws: Any, sources: list[Source]) -> None:
    ws.append(("來源代碼", "國家", "中文名稱", "原文名稱", "機關屬性", "語言", "IANA 時區", "首頁", "新聞入口", "RSS／Atom", "必要來源", "日期可省略", "卡片 selector", "日期 selector", "具名 parser", "健康基線觀察次數", "新鮮度門檻天數"))
    for source in sources:
        ws.append(
            (
                source.id,
                source.country,
                safe_excel_text(source.name_zh),
                safe_excel_text(source.name),
                safe_excel_text(source.institution_type),
                source.language,
                source.timezone,
                source.homepage,
                source.listing_url,
                "\n".join(source.feed_urls),
                "是" if source.critical else "否",
                "是" if source.date_optional else "否",
                "\n".join(source.card_selectors),
                "\n".join(source.date_selectors),
                source.parser_adapter,
                source.observation_runs,
                source.freshness_days,
            )
        )
        for column in (8, 9):
            ws.cell(ws.max_row, column).hyperlink = ws.cell(ws.max_row, column).value
            ws.cell(ws.max_row, column).style = "Hyperlink"
    _style_table(ws, widths=(22, 10, 28, 42, 18, 10, 22, 50, 64, 64, 12, 12, 40, 40, 20, 18, 18))


def _write_coverage(ws: Any, sources: list[Source]) -> None:
    source_names = {source.id: source.name_zh for source in sources}
    ws.append(("觀測議題", "法域", "主管／主要執行機關與研究生態系", "本程式監測來源代碼", "監測機關", "角色／機構性質", "官方機構／職掌證據網址", "最後查核日期"))
    for row in load_coverage(sources, validate_source_ids=False):
        ws.append(
            (
                row.topic,
                row.country,
                row.authorities,
                "；".join(row.source_ids),
                "；".join(source_names.get(source_id, source_id) for source_id in row.source_ids),
                row.roles,
                "\n".join(row.evidence_urls),
                row.verified_on,
            )
        )
        ws.cell(ws.max_row, 7).hyperlink = row.evidence_urls[0]
        ws.cell(ws.max_row, 7).style = "Hyperlink"
    _style_table(ws, widths=(48, 10, 76, 62, 76, 76, 76, 16))


def _style_table(ws: Any, widths: tuple[int, ...], max_row_height: float = 90) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        estimated_lines = 1
        for cell, width in zip(row, widths):
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            text = str(cell.value or "")
            estimated_lines = max(
                estimated_lines,
                sum(max(1, math.ceil(len(part) / max(8, width))) for part in text.splitlines() or [""]),
            )
        ws.row_dimensions[row[0].row].height = min(max_row_height, max(18, estimated_lines * 15))
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions


def _verify_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {
        "全部命中新聞", "CRA_CSA_NIS2_CER", "官方規範與執法", "研究智庫與公私協力",
        "來源健康狀態", "官方來源清單", "議題主管機關覆蓋",
    }
    missing = required - set(workbook.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")


def write_run_summary(
    output_path: str | Path,
    *,
    started_at: datetime,
    finished_at: datetime,
    since: datetime,
    until: datetime,
    articles: list[Article],
    statuses: list[SourceStatus],
    run_id: str = "",
    translation_report: TranslationReport | None = None,
    period: PeriodSelection | None = None,
    discovered_article_count: int | None = None,
    config_path: str | Path | None = None,
    run_profile: dict[str, Any] | None = None,
    artifact_names: list[str] | None = None,
) -> Path:
    workbook_path = Path(output_path).expanduser().resolve()
    summary_path = workbook_path.with_suffix(".run.json")
    translation = translation_report or TranslationReport(0, 0)
    hard_failure = any(
        item.critical and (item.fetch_status == "failed" or item.health_status == "degraded")
        for item in statuses
    )
    attention = any(
        item.fetch_status == "failed" or item.health_status in {"attention", "degraded"}
        for item in statuses
    ) or translation.success_rate < 0.95
    period_payload = period.as_dict() if period else {
        "mode": "fixed",
        "timezone": "UTC",
        "normalized_since": since.date().isoformat(),
        "normalized_until": (until - timedelta(microseconds=1)).date().isoformat(),
        "period_start_utc": since.isoformat(),
        "period_end_exclusive_utc": until.isoformat(),
    }
    discovered_count = len(articles) if discovered_article_count is None else discovered_article_count
    source_config_path = Path(config_path).expanduser().resolve() if config_path else default_sources_path().resolve()
    artifacts = artifact_names or [workbook_path.name, summary_path.name]
    payload = {
        "schema_version": 5,
        "program_version": _program_version(),
        "git_sha": _git_sha(),
        "python_version": platform.python_version(),
        "source_config_sha256": _file_sha256(source_config_path),
        "run_profile": run_profile or {},
        "run_id": run_id,
        "started_at": started_at.isoformat(),
        "finished_at": finished_at.isoformat(),
        "period_start": since.isoformat(),
        "period_end_exclusive": until.isoformat(),
        "period": period_payload,
        "output_file": str(workbook_path),
        "artifacts": artifacts,
        "article_count": len(articles),
        "discovered_article_count": discovered_count,
        "deduplicated_article_count": discovered_count - len(articles),
        "status": "degraded" if hard_failure else ("attention" if attention else "complete"),
        "translation": {
            "enabled": translation.enabled,
            "total": translation.total,
            "translated": translation.translated,
            "success_rate": round(translation.success_rate, 4),
            "failed_titles": list(translation.failed_titles),
            "cache_hits": translation.cache_hits,
            "providers": [item.__dict__ for item in translation.provider_stats],
        },
        "source_summary": {
            "total": len(statuses),
            "fetch_ok": sum(item.fetch_status == "ok" for item in statuses),
            "fetch_failed": sum(item.fetch_status == "failed" for item in statuses),
            "healthy": sum(item.health_status == "healthy" for item in statuses),
            "attention": sum(item.health_status == "attention" for item in statuses),
            "degraded": sum(item.health_status == "degraded" for item in statuses),
            "critical_degraded": sum(item.critical and item.health_status == "degraded" for item in statuses),
            "with_health_alerts": sum(bool(item.health_alerts) for item in statuses),
            "with_content_hits": sum(item.relevant_count > 0 for item in statuses),
            "invalid_date_count": sum(item.invalid_date_count for item in statuses),
            "timeout_count": sum(item.timeout_count for item in statuses),
            "budget_exhausted_count": sum(item.budget_exhausted for item in statuses),
            "request_count": sum(item.request_count for item in statuses),
            "bytes_downloaded": sum(item.bytes_downloaded for item in statuses),
            "retry_count": sum(item.retry_count for item in statuses),
        },
        "date_quality": {
            "dated_articles": sum(item.published_at is not None for item in articles),
            "undated_articles": sum(item.published_at is None for item in articles),
            "conflict_count": sum(item.date_conflict for item in articles),
            "by_source": _count_values(item.date_source or "unknown" for item in articles),
            "by_confidence": _count_values(item.date_confidence or "unknown" for item in articles),
        },
        "sources": [status.__dict__ for status in statuses],
    }
    _atomic_write_text(summary_path, json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    return summary_path


def _excel_summary(value: str) -> str:
    if len(value) <= EXCEL_SUMMARY_LIMIT:
        return value
    return value[: EXCEL_SUMMARY_LIMIT - 1].rstrip() + "…"


def _atomic_write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            stream.write(value)
            stream.flush()
            os.fsync(stream.fileno())
        temporary.replace(path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise


def _program_version() -> str:
    try:
        return version("eu-cyber-news-scraper")
    except PackageNotFoundError:
        return "0+unknown"


def _git_sha() -> str:
    try:
        return subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=Path(__file__).resolve().parent,
            check=True,
            capture_output=True,
            text=True,
            timeout=3,
        ).stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return ""


def _file_sha256(path: Path) -> str:
    try:
        return sha256(path.read_bytes()).hexdigest()
    except OSError:
        return ""


def _count_values(values: Iterable[str]) -> dict[str, int]:
    result: dict[str, int] = {}
    for value in values:
        result[value] = result.get(value, 0) + 1
    return result
