from __future__ import annotations

import json
import shutil
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .schema_validation import validate_schema_payload


class ArtifactBundle:
    """Stage and publish a run as a manifest-last, rollback-capable bundle."""

    def __init__(self, output: Path, run_id: str) -> None:
        self.output = output.expanduser().resolve()
        self.run_id = run_id
        self.root = self.output.parent / f".{self.output.stem}.{run_id}.staging"
        self.backup = self.root / "backup"

    def __enter__(self) -> ArtifactBundle:
        self.output.parent.mkdir(parents=True, exist_ok=True)
        shutil.rmtree(self.root, ignore_errors=True)
        self.root.mkdir()
        return self

    def __exit__(self, *_args: object) -> None:
        shutil.rmtree(self.root, ignore_errors=True)

    def staged(self, final_path: Path) -> Path:
        return self.root / final_path.name

    def publish(self, files: list[tuple[Path, Path]], *, manifest: Path) -> None:
        ordered = [pair for pair in files if pair[1] != manifest]
        ordered.extend(pair for pair in files if pair[1] == manifest)
        if not ordered or ordered[-1][1] != manifest:
            raise ValueError("artifact manifest must be published last")
        self.backup.mkdir()
        published: list[Path] = []
        backed_up: list[tuple[Path, Path]] = []
        try:
            for staged, final in ordered:
                if not staged.is_file():
                    raise FileNotFoundError(f"staged artifact is missing: {staged.name}")
                if final.exists():
                    backup = self.backup / final.name
                    final.replace(backup)
                    backed_up.append((backup, final))
                staged.replace(final)
                published.append(final)
        except BaseException:
            for path in reversed(published):
                path.unlink(missing_ok=True)
            for backup, final in reversed(backed_up):
                backup.replace(final)
            raise


def verify_artifact_bundle(
    workbook_path: Path,
    jsonl_path: Path | None,
    summary_path: Path,
    *,
    run_id: str,
    article_count: int,
) -> None:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        workbook_count = max(0, int(workbook["全部命中新聞"].max_row) - 1)
        metadata = {
            str(key): value
            for key, value in workbook["_run_metadata"].iter_rows(min_row=2, max_col=2, values_only=True)
            if key
        }
    finally:
        workbook.close()
    if workbook_count != article_count:
        raise ValueError(f"workbook article count mismatch: {workbook_count} != {article_count}")
    if metadata.get("schema_version") != 6 or metadata.get("run_id") != run_id:
        raise ValueError("workbook run_id or schema_version mismatch")
    if metadata.get("article_count") != article_count:
        raise ValueError("workbook metadata article count mismatch")

    if jsonl_path is not None:
        rows = [json.loads(line) for line in jsonl_path.read_text(encoding="utf-8").splitlines() if line]
        if len(rows) != article_count:
            raise ValueError(f"JSONL article count mismatch: {len(rows)} != {article_count}")
        if any(row.get("run_id") != run_id or row.get("schema_version") != 6 for row in rows):
            raise ValueError("JSONL run_id or schema_version mismatch")
        for row in rows:
            validate_schema_payload(row, "article-v6.schema.json")

    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    if summary.get("run_id") != run_id or summary.get("schema_version") != 6:
        raise ValueError("run manifest identity mismatch")
    validate_schema_payload(summary, "run-v6.schema.json")
    if summary.get("article_count") != article_count:
        raise ValueError("run manifest article count mismatch")
    if not summary.get("artifact_bundle_complete"):
        raise ValueError("run manifest does not mark the artifact bundle complete")
    manifest_period = summary.get("period", {})
    workbook_period = {
        key: metadata.get(key, "")
        for key in ("normalized_since", "normalized_until", "period_start_utc", "period_end_exclusive_utc")
    }
    if any(manifest_period.get(key, "") != value for key, value in workbook_period.items()):
        raise ValueError("workbook and run manifest period mismatch")
    artifacts = summary.get("artifact_manifest", [])
    if not artifacts or any(not row.get("complete") or not row.get("relative_path") for row in artifacts):
        raise ValueError("run manifest artifact entries are incomplete")
    actual_paths = [workbook_path, *([jsonl_path] if jsonl_path is not None else [])]
    actual_by_name = {path.name: path for path in actual_paths}
    declared_by_name = {str(row.get("name", "")): row for row in artifacts}
    if (
        len(artifacts) != len(actual_paths)
        or len(declared_by_name) != len(artifacts)
        or set(declared_by_name) != set(actual_by_name)
    ):
        raise ValueError("run manifest artifact set mismatch")
    for name, row in declared_by_name.items():
        relative_path = str(row["relative_path"])
        if relative_path != name or Path(relative_path).name != relative_path:
            raise ValueError(f"unsafe artifact relative path: {relative_path}")
        path = actual_by_name[name]
        expected = {
            "sha256": _file_sha256(path),
            "size_bytes": path.stat().st_size,
            "row_count": _artifact_row_count(path),
        }
        for field, actual in expected.items():
            if row.get(field) != actual:
                raise ValueError(f"artifact {field} mismatch for {name}")


def _file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _artifact_row_count(path: Path) -> int:
    if path.suffix == ".jsonl":
        return sum(bool(line.strip()) for line in path.read_text(encoding="utf-8").splitlines())
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return max(0, int(workbook["全部命中新聞"].max_row) - 1)
    finally:
        workbook.close()
