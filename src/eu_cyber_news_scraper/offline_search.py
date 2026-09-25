"""Search retained pre-filter official articles without network requests."""
from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from contextlib import closing
from datetime import datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook

from .config import load_sources_and_registry
from .exporter import safe_excel_text
from .models import Article
from .periods import resolve_period
from .ranking import is_hybrid_relevant, rank_articles
from .topic_profile import apply_profile, load_profile
from .topics import classify_article


def _index(connection: sqlite3.Connection, paths: list[Path]) -> None:
    connection.executescript("""
        CREATE TABLE IF NOT EXISTS files(path TEXT PRIMARY KEY, size INTEGER, mtime INTEGER);
        CREATE TABLE IF NOT EXISTS articles(file_path TEXT, url TEXT, published_at TEXT, source_id TEXT,
            payload TEXT, PRIMARY KEY(file_path, url));
        CREATE INDEX IF NOT EXISTS articles_date_source ON articles(published_at, source_id);
        CREATE TABLE IF NOT EXISTS search_cache(key TEXT PRIMARY KEY, workbook BLOB);
    """)
    present = {str(path.resolve()) for path in paths}
    for (old_path,) in connection.execute("SELECT path FROM files").fetchall():
        if old_path not in present:
            connection.execute("DELETE FROM articles WHERE file_path=?", (old_path,))
            connection.execute("DELETE FROM files WHERE path=?", (old_path,))
    for path in paths:
        stat = path.stat()
        key = str(path.resolve())
        previous = connection.execute("SELECT size,mtime FROM files WHERE path=?", (key,)).fetchone()
        if previous == (stat.st_size, stat.st_mtime_ns):
            continue
        connection.execute("DELETE FROM articles WHERE file_path=?", (key,))
        with path.open(encoding="utf-8") as stream:
            for line in stream:
                if not line.strip():
                    continue
                row = json.loads(line)
                connection.execute(
                    "INSERT OR REPLACE INTO articles VALUES (?,?,?,?,?)",
                    (key, row["url"], row.get("published_at"), row["source_id"], line.strip()),
                )
        connection.execute("INSERT OR REPLACE INTO files VALUES (?,?,?)", (key, stat.st_size, stat.st_mtime_ns))
    connection.commit()


def _article(row: dict[str, object]) -> Article:
    published = row.get("published_at")
    article = Article(
        source_id=str(row["source_id"]), country=str(row["country"]),
        source_name=str(row["source_name"]), institution_type=str(row["institution_type"]),
        language=str(row["language"]), title=str(row["title"]), url=str(row["url"]),
        published_at=datetime.fromisoformat(str(published)) if published else None,
        summary=str(row.get("summary") or ""),
    )
    article.published_date_local = str(row.get("published_date_local") or "")
    article.date_confidence = str(row.get("date_confidence") or "")
    article.title_zh_tw = str(row.get("title_zh_tw") or "")
    return article


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="從保留的官方文章 JSONL 離線重查主題，預設輸出 Excel。")
    parser.add_argument("--corpus-dir", type=Path, required=True)
    parser.add_argument("--topics-json")
    parser.add_argument("--since")
    parser.add_argument("--until")
    parser.add_argument("--days", type=int)
    parser.add_argument("--source", action="append")
    parser.add_argument("--topic", action="append", help="JSON 中的完整主題名稱；可重複。")
    parser.add_argument("--output", type=Path, default=Path("eu-search.xlsx"))
    args = parser.parse_args(argv)
    try:
        profile = load_profile(args.topics_json)
        settings = profile.manual
        if args.since is not None or args.until is not None:
            period = resolve_period(args.since, args.until, None)
        elif args.days is not None:
            period = resolve_period(None, None, args.days)
        elif settings.get("since") or settings.get("until"):
            period = resolve_period(settings.get("since"), settings.get("until"), None)
        else:
            period = resolve_period(None, None, settings.get("days", 14))
    except ValueError as exc:
        raise SystemExit(f"[error] {exc}") from exc
    if args.topic and set(args.topic) - profile.names:
        raise SystemExit(f"[error] 未知主題：{', '.join(sorted(set(args.topic) - profile.names))}")
    paths = sorted(args.corpus_dir.glob("*.corpus.jsonl"))
    if not paths:
        raise SystemExit("[error] 找不到篩選前官方文章 *.corpus.jsonl")
    args.corpus_dir.mkdir(parents=True, exist_ok=True)
    sources, registry = load_sources_and_registry()
    wanted_sources = set(args.source or [source.id for source in sources if not source.is_paused(period.end_date)])
    database = args.corpus_dir / ".offline-search.sqlite3"
    with closing(sqlite3.connect(database)) as connection:
        _index(connection, paths)
        # Coverage is inferred from the source run periods, not merely dates of published articles.
        intervals: list[tuple[str, str]] = []
        for path in paths:
            manifest = path.with_name(path.name.replace(".corpus.jsonl", ".run.json"))
            if manifest.exists():
                payload = json.loads(manifest.read_text(encoding="utf-8"))
                window = payload.get("period", {})
                successful = {row.get("source_id") for row in payload.get("sources", [])
                              if row.get("fetch_status") == "ok" and row.get("parse_status") not in {"empty", "failed", "not_run"}}
                if payload.get("artifact_bundle_complete") and payload.get("quality_gate", {}).get("passed") and wanted_sources <= successful:
                    intervals.append((window.get("period_start_utc", ""), window.get("period_end_exclusive_utc", "")))
        end = ""
        for start, stop in sorted(intervals):
            if start <= (end or period.since.isoformat()):
                end = max(end, stop)
            elif end < period.until.isoformat():
                break
        covered = bool(intervals) and intervals[0][0] <= period.since.isoformat() and end >= period.until.isoformat()
        if not covered:
            print("[warning] 保存資料未涵蓋完整查詢期間；結果僅為部分搜尋。", file=sys.stderr)
        fingerprint = sha256(repr((
            profile.hash, period.since.isoformat(), period.until.isoformat(), args.source, args.topic,
            [(str(path), path.stat().st_size, path.stat().st_mtime_ns) for path in paths],
            [(str(path), path.stat().st_mtime_ns) for path in (item.with_name(item.name.replace(".corpus.jsonl", ".run.json")) for item in paths) if path.exists()],
        )).encode()).hexdigest()
        cached = connection.execute("SELECT workbook FROM search_cache WHERE key=?", (fingerprint,)).fetchone()
        if cached:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_bytes(cached[0])
            print(f"{args.output}：沿用未變規則與文章索引；資料涵蓋{'完整' if covered else '不完整'}")
            return
        clause = "WHERE published_at>=? AND published_at<?"
        params: list[str] = [period.since.isoformat(), period.until.isoformat()]
        if args.source:
            clause += f" AND source_id IN ({','.join('?' for _ in args.source)})"
            params.extend(args.source)
        rows = connection.execute(f"SELECT payload FROM articles {clause}", params).fetchall()
    unique: dict[str, Article] = {}
    for (value,) in rows:
        article = _article(json.loads(value))
        unique[article.url] = article
    articles = list(unique.values())
    for article in articles:
        classify_article(article)
        apply_profile(article, profile)
    rank_articles(articles, registry, profile=profile)
    matches = sorted((item for item in articles if is_hybrid_relevant(item) and (not args.topic or set(args.topic) & set(item.matched_topics))), key=lambda item: item.published_at or period.since, reverse=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("查詢結果")
    sheet.append(("日期", "來源", "標題", "繁體標題", "主題", "關鍵詞", "責任機關", "官方網址"))
    for item in matches:
        sheet.append(tuple(safe_excel_text(value) for value in (
            item.published_date_local or (item.published_at.date().isoformat() if item.published_at else ""),
            item.source_name, item.title, item.title_zh_tw, "；".join(item.matched_topics),
            "；".join(item.matched_keywords), "；".join(item.responsibility_owner or ["未設定"]), item.url,
        )))
    info = workbook.create_sheet("查詢資訊")
    info.append(("規則 SHA-256", profile.hash))
    info.append(("查詢開始", period.start_date.isoformat()))
    info.append(("查詢結束（含當日）", period.end_date.isoformat()))
    info.append(("資料涵蓋完整", "是" if covered else "否"))
    info.append(("筆數", len(matches)))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    with closing(sqlite3.connect(database)) as connection:
        connection.execute("INSERT OR REPLACE INTO search_cache VALUES (?,?)", (fingerprint, args.output.read_bytes()))
        connection.commit()
    print(f"{args.output}：{len(matches)} 筆；資料涵蓋{'完整' if covered else '不完整'}；規則 {profile.hash[:12]}")
