from datetime import datetime, timezone

from openpyxl import load_workbook

from eu_cyber_news_scraper.exporter import EXCEL_SUMMARY_LIMIT, export_workbook, safe_excel_text, write_run_summary
from eu_cyber_news_scraper.models import Article, SourceStatus


def test_exporter_creates_required_sheets_and_summary(tmp_path):
    article = Article(
        source_id="eu",
        country="EU",
        source_name="歐盟機關",
        institution_type="歐盟行政機關",
        language="en",
        title="Cyber Resilience Act guidance",
        title_zh_tw="《網路韌性法》指引",
        url="https://example.eu/news/cra",
        published_at=datetime(2026, 5, 1, tzinfo=timezone.utc),
        matched_topics=["產品資安、漏洞揭露義務、SBOM"],
        matched_keywords=["cyber resilience act"],
        relevance_score=4,
        confidence_level="高",
        summary="x" * 700,
        alternate_urls=["https://mirror.example.eu/news/cra"],
    )
    status = SourceStatus("eu", "歐盟機關", "EU", True, True, "feed", 1, 1, 0.2)
    # Coverage validation requires the production source inventory.
    from eu_cyber_news_scraper.config import load_sources

    sources = list(load_sources())
    path = export_workbook([article], [status], sources, tmp_path / "result.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "全部命中新聞",
        "CRA_CSA_NIS2_CER",
        "官方規範與執法",
        "研究智庫與公私協力",
        "來源健康狀態",
        "官方來源清單",
        "議題主管機關覆蓋",
    ]
    assert workbook["議題主管機關覆蓋"].max_row == 61
    assert workbook["CRA_CSA_NIS2_CER"].max_row == 2
    assert workbook["全部命中新聞"]["Q2"].value == "《網路韌性法》指引"
    assert workbook["全部命中新聞"]["A2"].fill.fgColor.rgb == "00FFC000"
    assert workbook["CRA_CSA_NIS2_CER"]["A2"].fill.fill_type is None
    assert len(workbook["全部命中新聞"]["R2"].value) == EXCEL_SUMMARY_LIMIT
    assert workbook["全部命中新聞"]["X2"].value == "https://mirror.example.eu/news/cra"
    assert workbook["全部命中新聞"].row_dimensions[2].height == 90

    now = datetime.now(timezone.utc)
    summary = write_run_summary(
        path,
        started_at=now,
        finished_at=now,
        since=now,
        until=now,
        articles=[article],
        statuses=[status],
    )
    assert '"status": "complete"' in summary.read_text(encoding="utf-8")
    assert '"source_summary"' in summary.read_text(encoding="utf-8")
    assert '"schema_version": 5' in summary.read_text(encoding="utf-8")


def test_excel_formula_text_is_escaped():
    assert safe_excel_text("=HYPERLINK(\"bad\")").startswith("'")
    assert safe_excel_text("ordinary title") == "ordinary title"
