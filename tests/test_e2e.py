import json
from types import SimpleNamespace

from openpyxl import load_workbook

from eu_cyber_news_scraper.cli import _run_pipeline
from eu_cyber_news_scraper.models import Source
from eu_cyber_news_scraper.periods import resolve_period


class FakeHttpClient:
    def __init__(self, timeout: int):
        self.payload = b""

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        return None

    async def get(self, url: str, *, stats=None):
        import httpx

        return httpx.Response(
            200,
            content=(
                b"<?xml version='1.0'?><rss><channel><item><title>NIS2 security guidance</title>"
                b"<link>https://agency.example/news/nis2</link>"
                b"<pubDate>Fri, 01 May 2026 08:00:00 GMT</pubDate>"
                b"<description>NIS2 risk management requirements</description></item></channel></rss>"
            ),
            request=httpx.Request("GET", url),
        )


def test_offline_pipeline_writes_atomic_artifacts_and_quality_metadata(monkeypatch, tmp_path):
    monkeypatch.setattr("eu_cyber_news_scraper.cli.HttpClient", FakeHttpClient)
    monkeypatch.setattr(
        "eu_cyber_news_scraper.translation._translate_titles",
        lambda titles: {title: "NIS2 資安指引" for title in titles},
    )
    monkeypatch.setenv("EU_CYBER_NEWS_TRANSLATION_CACHE", str(tmp_path / "translations.json"))
    source = Source(
        "test", "EU", "測試主管機關", "Test Authority", "國家主管機關", "en",
        "https://agency.example", "https://agency.example/news", feed_urls=("https://agency.example/rss",),
        allow_domains=("agency.example",), include_patterns=(r"/news/.+",), critical=True,
    )
    args = SimpleNamespace(
        timeout=1, workers=1, all=False, include_undated=False, no_detail=True,
        topic=None, jsonl=True, fail_on_degraded=False, source_budget=10, state_dir=str(tmp_path / ".state"),
        health_write="always",
    )
    output = tmp_path / "result.xlsx"
    period = resolve_period("1150501", "1150502", None)
    _run_pipeline(
        args, [source], period.since, period.until, output, "run-e2e", period=period,
    )
    workbook = load_workbook(output, read_only=True)
    assert "官方規範與執法" in workbook.sheetnames
    summary = json.loads(output.with_suffix(".run.json").read_text(encoding="utf-8"))
    assert summary["run_id"] == "run-e2e"
    assert summary["translation"]["success_rate"] == 1
    assert summary["schema_version"] == 5
    assert summary["period"]["raw_since"] == "1150501"
    assert summary["period"]["since_calendar"] == "roc"
    assert (tmp_path / ".state" / ".source-health.json").exists()
    assert output.with_suffix(".jsonl").exists()
    row = json.loads(output.with_suffix(".jsonl").read_text(encoding="utf-8").splitlines()[0])
    assert row["schema_version"] == 5
    assert row["run_id"] == "run-e2e"
    assert row["date_source"] == "feed-published"
