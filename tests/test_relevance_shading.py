from dataclasses import asdict, replace
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook

from eu_cyber_news_scraper.config import load_sources
from eu_cyber_news_scraper.exporter import ARTICLE_HEADERS, _write_articles, export_workbook, relevance_fill
from eu_cyber_news_scraper.models import Article, SourceStatus

ARTICLE_VIEWS = (
    "全部命中新聞", "CRA_CSA_NIS2_CER", "官方規範與執法", "研究智庫與公私協力",
)
SHADED_HEADERS = {
    "新聞標題", "繁體中文標題", "摘要", "觀測主題", "命中關鍵字", "關聯分數",
    "Boolean 分數", "BM25 分數", "BM25 主題分數", "命中同義詞", "可信度",
}


def _article(**changes):
    article = Article(
        source_id="eu",
        country="EU",
        source_name="歐盟機關",
        institution_type="歐盟行政機關",
        language="en",
        title="Cyber Resilience Act guidance",
        title_zh_tw="《網路韌性法》指引",
        summary="Cyber resilience guidance for products.",
        url="https://example.eu/news/cra",
        alternate_urls=["https://mirror.example.eu/news/cra"],
        published_at=datetime(2026, 5, 1, tzinfo=timezone.utc),
        published_at_raw="1 May 2026",
        published_date_local="2026-05-01",
        date_confidence="low",
        matched_topics=["產品資安、漏洞揭露義務、SBOM"],
        matched_keywords=["cyber resilience act"],
        matched_synonyms=["CRA"],
        relevance_score=4,
        boolean_score=4,
        bm25_score=2.125,
        bm25_topic_scores={"產品資安、漏洞揭露義務、SBOM": 2.125},
        confidence_level="高",
        authority_level="主管／監理機關",
    )
    return replace(article, **changes)


@pytest.fixture(params=[("高", "FFD966", 4), ("中", "FFE699", 2), ("低", "FFF2CC", 1)])
def shaded_workbook(request, tmp_path):
    level, color, score = request.param
    articles = [
        _article(confidence_level=level, relevance_score=score, boolean_score=score, matched_synonyms=[]),
        _article(
            confidence_level=level, relevance_score=score, boolean_score=score,
            source_id="research", authority_level="研究／支援機構",
            url="https://example.eu/research/cra",
            matched_keywords=[],
        ),
    ]
    before = [asdict(article) for article in articles]
    for article in articles:
        assert relevance_fill(article).fgColor.rgb == f"00{color}"
    sources = list(load_sources())
    status = SourceStatus("eu", "歐盟機關", "EU", True, True, "feed", 2, 2, 0.2)
    path = export_workbook(articles, [status], sources, tmp_path / "shading.xlsx", run_id="shading-check")
    assert [asdict(article) for article in articles] == before
    workbook = load_workbook(path)
    yield workbook, articles, color, sources
    workbook.close()


def test_exact_yellow_palette_and_target_columns_in_all_article_views(shaded_workbook):
    workbook, articles, color, _ = shaded_workbook
    for name in ARTICLE_VIEWS:
        ws = workbook[name]
        assert tuple(cell.value for cell in ws[1]) == ARTICLE_HEADERS
        assert ws.max_column == 33
        expected = articles if name in ARTICLE_VIEWS[:2] else [articles[0 if name == ARTICLE_VIEWS[2] else 1]]
        assert ws.max_row == len(expected) + 1
        for row in ws.iter_rows(min_row=2):
            for header, cell in zip(ARTICLE_HEADERS, row, strict=True):
                if header in SHADED_HEADERS:
                    assert cell.fill.fill_type == "solid", (name, header)
                    assert cell.fill.fgColor.rgb == f"00{color}", (name, header)
                else:
                    assert cell.fill.fill_type is None, (name, header)


@pytest.mark.parametrize(
    "changes,expected_color",
    [
        ({"matched_keywords": [], "matched_synonyms": ["CRA"]}, "FFD966"),
        ({"matched_keywords": [], "matched_synonyms": []}, None),
        ({"matched_topics": []}, None),
        ({"confidence_level": "未命中"}, None),
        ({"confidence_level": ""}, None),
        ({"confidence_level": "unknown"}, None),
        ({"confidence_level": "低", "date_confidence": "high"}, "FFF2CC"),
        ({"confidence_level": "", "date_confidence": "high"}, None),
    ],
)
def test_requires_topic_and_keyword_or_synonym_and_known_relevance(changes, expected_color):
    workbook = Workbook()
    ws = workbook.active
    article = _article(**changes)
    fill = relevance_fill(article)
    if expected_color:
        assert fill.fill_type == "solid"
        assert fill.fgColor.rgb == f"00{expected_color}"
    else:
        assert fill is None
    _write_articles(ws, [article])
    for header, cell in zip(ARTICLE_HEADERS, ws[2], strict=True):
        if expected_color and header in SHADED_HEADERS:
            assert cell.fill.fill_type == "solid"
            assert cell.fill.fgColor.rgb == f"00{expected_color}"
        else:
            assert cell.fill.fill_type is None


def test_layout_headers_counts_links_and_values_are_preserved(shaded_workbook):
    workbook, articles, _, _ = shaded_workbook
    for name in ARTICLE_VIEWS:
        ws = workbook[name]
        assert ws.freeze_panes == "C2"
        assert ws.sheet_view.showGridLines is False
        assert ws.auto_filter.ref == f"A1:AG{ws.max_row}"
        assert ws.column_dimensions["P"].width == 56
        assert ws.column_dimensions["Q"].width == 56
        assert ws.column_dimensions["R"].width == 72
        assert ws.column_dimensions["AD"].width == 64
        for cell in ws[1]:
            assert cell.fill.fgColor.rgb == "001F4E78"
            assert cell.font.bold is True
            assert cell.font.color.rgb == "00FFFFFF"
            assert cell.alignment.horizontal == "center"
        by_url = {article.url: article for article in articles}
        for index, row in enumerate(ws.iter_rows(min_row=2), 1):
            cells = dict(zip(ARTICLE_HEADERS, row, strict=True))
            link = cells["官方原文"]
            article = by_url[link.value]
            assert link.hyperlink.target == article.url
            assert link.style == "Hyperlink"
            assert link.font.color.type == "theme"
            assert link.font.color.theme == 10
            assert cells["編號"].value == index
            assert cells["新聞標題"].value == article.title
            assert cells["繁體中文標題"].value == article.title_zh_tw
            assert cells["摘要"].value == article.summary
            assert cells["觀測主題"].value == "；".join(article.matched_topics)
            assert cells["命中關鍵字"].value == ("；".join(article.matched_keywords) or None)
            assert cells["命中同義詞"].value == ("；".join(article.matched_synonyms) or None)
            assert cells["關聯分數"].value == article.relevance_score
            assert cells["Boolean 分數"].value == article.boolean_score
            assert cells["BM25 分數"].value == article.bm25_score
            assert cells["可信度"].value == article.confidence_level
            assert cells["日期信心"].value == article.date_confidence
            assert cells["其他官方連結"].value == article.alternate_urls[0]
            assert all(cell.alignment.wrap_text is True for cell in row)
            assert all(cell.alignment.vertical == "top" for cell in row)
            assert 18 <= ws.row_dimensions[row[0].row].height <= 90
    metadata = dict(workbook["_run_metadata"].values)
    assert metadata["schema_version"] == 6
    assert metadata["article_count"] == 2
    assert metadata["run_id"] == "shading-check"


def test_source_status_coverage_and_settings_body_remain_neutral(shaded_workbook):
    workbook, _, _, sources = shaded_workbook
    for name in ("來源健康狀態", "官方來源清單", "議題主管機關覆蓋", "篩選設定", "_run_metadata"):
        assert all(cell.fill.fill_type is None for row in workbook[name].iter_rows(min_row=2) for cell in row)
    source_sheet = workbook["官方來源清單"]
    assert source_sheet.max_row == len(sources) + 1
    assert workbook["來源健康狀態"].max_row == 2
    assert source_sheet["H2"].hyperlink.target == sources[0].homepage
    assert source_sheet["I2"].hyperlink.target == sources[0].listing_url
    assert source_sheet["H2"].style == "Hyperlink"
    settings = dict(workbook["篩選設定"].values)
    assert settings["filter_method"] == "boolean_then_bm25"
    assert "confidence_level" in settings["relevance_shading_basis"]
    assert "與日期信心無關" in settings["relevance_shading_basis"]
    assert settings["relevance_shading_high"] == "高：FFD966"
    assert settings["relevance_shading_medium"] == "中：FFE699"
    assert settings["relevance_shading_low"] == "低：FFF2CC"
    assert set(settings["relevance_shading_columns"].split("；")) == SHADED_HEADERS


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_shading_preserves_formula_injection_protection(prefix, tmp_path):
    text = f"{prefix}formula\btext"
    article = _article(title=text, title_zh_tw=text, summary=text)
    path = export_workbook([article], [], list(load_sources()), tmp_path / "escaped.xlsx")
    workbook = load_workbook(path)
    for name in ARTICLE_VIEWS[:3]:
        ws = workbook[name]
        for column in (16, 17, 18):
            cell = ws.cell(2, column)
            assert cell.value == f"'{prefix}formulatext"
            assert cell.data_type == "s"
            assert cell.fill.fgColor.rgb == "00FFD966"
    workbook.close()
